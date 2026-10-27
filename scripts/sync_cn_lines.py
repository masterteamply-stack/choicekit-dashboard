#!/usr/bin/env python3
"""중국차 오더 리포트(const CNORD)를 '오더라인' 시트 한 장으로 통째로 다시 만든다.

  python3 scripts/sync_cn_lines.py --lines "중국차 오더 월별 이력 2026_01-10.xlsx"
  python3 scripts/sync_cn_lines.py --lines <파일> --check     # 쓰지 않고 수치만

쓰는 시트는 '오더라인_2026' 이고 다음 열이 있어야 한다.
  월 · 매칭 코드 · 중국 브랜드 · 권역 (대시보드 기준) · 정가(USD)
  ② 중국전용 · ③ 신규 (237 기준) · ③b 신규 (중국전용 한정)   ← 전부 Y/N

왜 월 export 가 아니라 이 시트인가
  월 export 는 그 달에 받은 것만 담기는데 전표가 빠지는 수가 있다. 2026-10 export 는
  CTR EUROPE 판매문서 4건(802887·802888·802892·802893, 커버 품번 $17,584)이 통째로
  빠져 있었다. 오더라인 시트는 원장 3개년을 다 보고 만든 것이라 이쪽을 정본으로 쓴다.

신규(③) 두 가지를 같이 담는다
  it[2] = ③  중국 전용 품번만 — 지금까지 대시보드가 쓰던 정의
  it[6] = ③+ 공용 품번까지 — 237 코드 전체에 같은 규칙을 적용한 것
"""
import json, os, re, sys, collections

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HTML = os.path.join(ROOT, 'promo_dashboard_v10_6.html')
NEED = ['월', '매칭 코드', '중국 브랜드', '권역 (대시보드 기준)', '정가(USD)',
        '② 중국전용', '③ 신규 (237 기준)', '③b 신규 (중국전용 한정)']

if '--lines' not in sys.argv:
    print('--lines <오더라인 xlsx> 가 필요합니다.'); sys.exit(1)
import openpyxl
wb = openpyxl.load_workbook(sys.argv[sys.argv.index('--lines') + 1], data_only=True)
sh = next((s for s in wb.sheetnames if s.startswith('오더라인')), None)
if not sh:
    print("'오더라인…' 시트를 찾지 못했습니다: " + ', '.join(wb.sheetnames)); sys.exit(1)
rows = list(wb[sh].iter_rows(values_only=True))
H = {str(h).strip(): i for i, h in enumerate(rows[0]) if h is not None}
miss = [k for k in NEED if k not in H]
if miss:
    print('시트에 없는 열: ' + ', '.join(miss)); sys.exit(1)
Y = lambda v: str(v or '').strip().upper() == 'Y'

html = open(HTML, encoding='utf-8', newline='').read()
mC = re.search(r'const CNORD=(\{.*?\});\n', html, re.S)
if not mC:
    print('CNORD 블록을 찾지 못했습니다.'); sys.exit(1)
CN = json.loads(mC.group(1))
old = {it[0]: it for it in CN['items']}

per = collections.defaultdict(float)        # (code, region, month) -> usd
meta = {}
months, regions = set(), []
for r in rows[1:]:
    c = str(r[H['매칭 코드']] or '').strip()
    if not c: continue
    mm = '%02d' % int(r[H['월']])
    rg = str(r[H['권역 (대시보드 기준)']] or '').strip()
    v = r[H['정가(USD)']]
    v = float(v) if isinstance(v, (int, float)) else 0.0
    months.add(mm)
    if rg and rg not in regions: regions.append(rg)
    per[(c, rg, mm)] += v
    meta.setdefault(c, dict(cn=str(r[H['중국 브랜드']] or '').strip(),
                            only=Y(r[H['② 중국전용']]),
                            new237=Y(r[H['③ 신규 (237 기준)']]),
                            newonly=Y(r[H['③b 신규 (중국전용 한정)']])))

items = []
for c, m in sorted(meta.items()):
    o = old.get(c)
    st = o[4] if o else 'Mass Production'      # 상태는 기존 값을 잇는다
    g = {}
    for (cc, rg, mm), v in per.items():
        if cc == c and v: g.setdefault(rg, {})[mm] = round(v, 2)
    for rg in g: g[rg] = {k: g[rg][k] for k in sorted(g[rg])}
    items.append([c, 1 if m['only'] else 0, 1 if m['newonly'] else 0,
                  m['cn'] or (o[3] if o else ''), st, g, 1 if m['new237'] else 0])

ORD = ['CIS', 'ASIA', 'MENA', 'LATAM', 'NA', 'EU', 'OC']
regions = [r for r in ORD if r in regions] + [r for r in regions if r not in ORD]
CN.update(months=sorted(months), regions=regions, items=items,
          src_lines=os.path.basename(sys.argv[sys.argv.index('--lines') + 1]))
CN.pop('src_month', None)

S = lambda mm, pred: round(sum(v for (c, rg, m2), v in per.items() if m2 == mm and pred(c)), 2)
print(f"{sh} · {len(rows)-1:,}행 · 품번 {len(items)}개 · 월 {','.join(sorted(months))}")
print(f"{'월':>4}{'① 커버 전체':>15}{'② 중국 전용':>14}{'③ 중국전용 신규':>16}{'③+ 공용 포함':>14}")
for mm in sorted(months):
    print(f"{int(mm):>3}월{S(mm, lambda c: True):>15,.2f}{S(mm, lambda c: meta[c]['only']):>14,.2f}"
          f"{S(mm, lambda c: meta[c]['newonly']):>16,.2f}{S(mm, lambda c: meta[c]['new237']):>14,.2f}")
print(f"{'합계':>4}{S(None, lambda c: True) if False else sum(per.values()):>15,.2f}")

out = html[:mC.start()] + 'const CNORD=' + json.dumps(CN, ensure_ascii=False, separators=(',', ':')) + ';\n' + html[mC.end():]
changed = out != html
if '--check' in sys.argv:
    print('변경 있음' if changed else '변경 없음'); sys.exit(0)
if changed: open(HTML, 'w', encoding='utf-8', newline='').write(out)
print('promo_dashboard_v10_6.html ' + ('갱신됨' if changed else '변경 없음'))

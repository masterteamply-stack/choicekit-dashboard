#!/usr/bin/env python3
"""월 오더 export(SAP 수출 원장 · Plnt 1150)에서 캠페인별 그 달 오더를 뽑아
campaign_order_add.json 에 쌓고, 대시보드의 const D 캠페인 행에 더한다.

  python3 scripts/sync_campaign_month.py --xlsx export.XLSX
  python3 scripts/sync_campaign_month.py              # json 만으로 다시 반영
  python3 scripts/sync_campaign_month.py --check      # 쓰지 않고 변경 여부만

왜 캠페인 전부가 아니라 일부만 다루나
  이 export 는 Plnt 1150(본사 수출) 한 곳이다. Lordco(CTR A)·ASIA 일부(CTR V)처럼
  다른 법인을 거치는 고객은 통째로 빠져 있어, 그 캠페인에 더하면 실적이 깎인다.
  그래서 참여 고객이 모두 1150 을 거치는 캠페인만 COVERED 에 적어 둔다.

대상 판정은 campaign_target_skus.json 의 규칙을 그대로 쓴다
  match_by=sku  대상 SKU 목록 ∩ 그 캠페인 권역 ∩ (목록이 있으면) 대상 고객
  match_by=cust 대상 고객 ∩ 그 캠페인 권역
"""
import json, os, re, sys, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HTML = os.path.join(ROOT, 'promo_dashboard_v10_6.html')
ADD = os.path.join(ROOT, 'campaign_order_add.json')
TGT = os.path.join(ROOT, 'campaign_target_skus.json')

# 이 export(Plnt 1150) 하나로 그 달 실적이 온전히 잡히는 캠페인만 더한다.
COVERED = {'MENA|Portfolio Expansion Program'}
# 오더로 보지 않는 건 — 캠페인 주석에 적힌 예외
EXCLUDE = {'MENA|Portfolio Expansion Program': {'EURL TASSRIAA'}}
# export 의 판매처가 D.customers 에 없는 경우의 권역
REG_FIX = {'200505': 'CIS', '289087': 'EUROPE', '200716': 'MENA',
           '200666': 'NA', '200219': 'CIS', '200697': 'CIS'}
REG_MAP = {'Korea': 'KR'}

nrm = lambda s: re.sub(r'[\s.,()·\-]', '', str(s or '')).upper()


def read_export(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    head = next(it)
    ix = {str(h).strip(): i for i, h in enumerate(head) if h is not None}
    need = ['CTR Code', '판매처', '판매처(명)', '오더 수량', '정가(USD)', '생성일']
    miss = [k for k in need if k not in ix]
    if miss:
        print('export 에 없는 열: ' + ', '.join(miss)); sys.exit(1)
    out = []
    for r in it:
        if not r or not r[ix['CTR Code']]: continue
        d = r[ix['생성일']]
        d = d.date().isoformat() if isinstance(d, datetime.datetime) else str(d)[:10]
        out.append(dict(ctr=str(r[ix['CTR Code']]).strip(), cust=str(r[ix['판매처']] or '').strip(),
                        name=str(r[ix['판매처(명)']] or '').strip(),
                        qty=float(r[ix['오더 수량']] or 0), usd=float(r[ix['정가(USD)']] or 0), d=d))
    return out


html = open(HTML, encoding='utf-8', newline='').read()
mD = re.search(r'<script>const D = (\{.*?\});</script>', html, re.S)
if not mD:
    print('D 블록을 찾지 못했습니다.'); sys.exit(1)
D = json.loads(mD.group(1))
REG_OF = {c['k']: REG_MAP.get(c.get('r'), c.get('r')) for c in D.get('customers', [])}
CTRY = {c['k']: (c.get('c') or '') for c in D.get('customers', [])}
T = json.load(open(TGT, encoding='utf-8'))
add = json.load(open(ADD, encoding='utf-8')) if os.path.exists(ADD) else {'_readme':
    '월 오더 export 에서 뽑은 캠페인별 월 추가분. scripts/sync_campaign_month.py 로 만든다.', 'camps': {}}

if '--xlsx' in sys.argv:
    xl = sys.argv[sys.argv.index('--xlsx') + 1]
    rows = read_export(xl)
    month = max(r['d'] for r in rows)[:7]
    print(f"export {os.path.basename(xl)} · {len(rows):,}행 · {month} · 최종 오더일 {max(r['d'] for r in rows)}")
    for key in sorted(COVERED):
        rg, pg = key.split('|')
        t = T.get(pg)
        if not t:
            print(f'  {key} — campaign_target_skus.json 에 규칙 없음'); continue
        skus = set(t.get('target_skus') or [])
        custs = {nrm(x) for x in (t.get('target_customers') or [])}
        excl = {nrm(x) for x in EXCLUDE.get(key, set())}
        sel = [r for r in rows
               if (REG_FIX.get(r['cust']) or REG_OF.get(r['cust'])) == rg
               and (not custs or nrm(r['name']) in custs)
               and nrm(r['name']) not in excl
               and (not skus or r['ctr'] in skus)]
        per = {}
        for r in sel:
            v = per.setdefault(r['name'], dict(code=r['cust'], usd=0.0))
            v['usd'] += r['usd']
        e = dict(usd=round(sum(r['usd'] for r in sel), 2), qty=int(sum(r['qty'] for r in sel)),
                 lines=len(sel), n_sku=len({r['ctr'] for r in sel}),
                 cut=(max(r['d'] for r in sel) if sel else ''),
                 custs=sorted({r['name'] for r in sel}),
                 by_cust={k: dict(code=v['code'], usd=round(v['usd'], 2)) for k, v in sorted(per.items())},
                 src=os.path.basename(xl))
        add['camps'].setdefault(key, {})[month] = e
        print(f"  {key} · {month} · ${e['usd']:,.2f} · {e['n_sku']}종 · {e['qty']:,}개 · {', '.join(e['custs'])}")
    json.dump(add, open(ADD, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ── D 캠페인 행에 더하기 (base_* 를 남겨 두어 몇 번 돌려도 같은 결과)
touched = []
for c in D.get('campaigns', []):
    key = f"{c.get('region')}|{c.get('program')}"
    months = add['camps'].get(key)
    if not months: continue
    c.setdefault('base_a26', c['a26'])
    c.setdefault('base_cut', c['cut'])
    c.setdefault('base_mo_ord', c.get('mo_ord'))
    last = max(months)
    c['a26'] = round(c['base_a26'] + sum(v['usd'] for v in months.values()), 2)
    c['mo_ord'] = months[last]['usd']
    c['cut'] = months[last]['cut'] or c['base_cut']
    c['add_months'] = {m: v['usd'] for m, v in sorted(months.items())}
    c['add_sku'] = {m: v['n_sku'] for m, v in sorted(months.items())}
    touched.append((key, c['base_a26'], c['a26'], last, c['mo_ord']))
    # 그 달에 처음 오더한 고객은 캠페인 고객 목록(cust_delta)에도 올린다
    cd = D.setdefault('cust_delta', [])
    have = {r['name'] for r in cd if r.get('region') == c['region'] and r.get('program') == c['program']}
    tot_by = {}
    for v in months.values():
        for nm, x in (v.get('by_cust') or {}).items():
            t = tot_by.setdefault(nm, dict(code=x['code'], usd=0.0)); t['usd'] += x['usd']
    for nm, x in sorted(tot_by.items()):
        if nm in have:
            for r in cd:
                if r.get('region') == c['region'] and r.get('program') == c['program'] and r['name'] == nm:
                    r.setdefault('base_a26', r['a26']); r['a26'] = round(r['base_a26'] + x['usd'], 2)
        else:
            cd.append(dict(code=x['code'], name=nm, country=(CTRY.get(x['code']) or ''),
                           region=c['region'], program=c['program'], a25=0.0,
                           a26=round(x['usd'], 2), base_a26=0.0))
    c['n_cust'] = sum(1 for r in cd if r.get('region') == c['region'] and r.get('program') == c['program'])

out = html[:mD.start()] + '<script>const D = ' + json.dumps(D, ensure_ascii=False, separators=(',', ':')) \
      + ';</script>' + html[mD.end():]
changed = out != html
if '--check' in sys.argv:
    print('변경 있음' if changed else '변경 없음'); sys.exit(0)
if changed: open(HTML, 'w', encoding='utf-8', newline='').write(out)

print()
for key, b, a, m, mo in touched:
    print(f"{key:44} a26 {b:>12,.0f} → {a:>12,.0f}  ({m} +{mo:,.0f})")
print('promo_dashboard_v10_6.html ' + ('갱신됨' if changed else '변경 없음'))
